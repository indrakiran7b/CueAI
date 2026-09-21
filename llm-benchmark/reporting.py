"""CSV / Excel / chart reporting — never writes API keys."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils.dataframe import dataframe_to_rows  # noqa: E402

from config import (
    CATEGORY_SUMMARY_CSV,
    CHART_ACCURACY,
    CHART_AVG_LATENCY,
    CHART_AVG_TTFT,
    CHART_CATEGORY_ACCURACY,
    CHART_CATEGORY_LATENCY,
    CHART_CUEAI_LATENCY,
    CHART_CUEAI_TTFT,
    CHART_P90_TTFT,
    CHART_QUESTION_LATENCY,
    CHART_SUCCESS,
    CHART_TPS,
    EXCEL_PATH,
    MODEL_SUMMARY_CSV,
    QUESTION_SUMMARY_CSV,
    RAW_CSV,
    RESULTS_DIR,
)


def ensure_results_dir() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)


def _safe_label(name: str, max_len: int = 16) -> str:
    return name if len(name) <= max_len else name[: max_len - 1] + "..."


def save_raw_csv(df: pd.DataFrame, path: Path = RAW_CSV) -> Path:
    ensure_results_dir()
    df.to_csv(path, index=False)
    return path


def save_model_summary_csv(df: pd.DataFrame, path: Path = MODEL_SUMMARY_CSV) -> Path:
    ensure_results_dir()
    df.to_csv(path, index=False)
    return path


def save_question_summary_csv(
    df: pd.DataFrame, path: Path = QUESTION_SUMMARY_CSV
) -> Path:
    ensure_results_dir()
    df.to_csv(path, index=False)
    return path


def save_category_summary_csv(
    df: pd.DataFrame, path: Path = CATEGORY_SUMMARY_CSV
) -> Path:
    ensure_results_dir()
    df.to_csv(path, index=False)
    return path


# Back-compat
def save_summary_csv(df: pd.DataFrame, path: Path = MODEL_SUMMARY_CSV) -> Path:
    return save_model_summary_csv(df, path)


def _bar_chart(
    labels: list[str],
    values: list[float],
    title: str,
    ylabel: str,
    out_path: Path,
) -> Optional[Path]:
    if not labels or not values or len(labels) != len(values):
        return None
    fig, ax = plt.subplots(figsize=(10, 4.8))
    bars = ax.bar(range(len(values)), values, color="#2F6FED")
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    for bar, val in zip(bars, values):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height(),
            f"{val:.1f}",
            ha="center",
            va="bottom",
            fontsize=8,
        )
    fig.tight_layout()
    ensure_results_dir()
    fig.savefig(out_path, dpi=140)
    plt.close(fig)
    return out_path


def _model_labels(summary: pd.DataFrame) -> list[str]:
    return [
        f"{r.Provider}\n{_safe_label(str(r.Model))}" for r in summary.itertuples()
    ]


def generate_charts(
    model_summary: pd.DataFrame, raw: pd.DataFrame
) -> dict[str, Optional[Path]]:
    charts: dict[str, Optional[Path]] = {
        "avg_ttft": None,
        "p90_ttft": None,
        "avg_latency": None,
        "tps": None,
        "success": None,
        "accuracy": None,
        "cueai_ttft": None,
        "cueai_latency": None,
        "category_latency": None,
        "category_accuracy": None,
        "question_latency": None,
    }
    if model_summary.empty:
        return charts

    labels = _model_labels(model_summary)

    def col_vals(col: str, scale: float = 1.0) -> list[float]:
        out: list[float] = []
        for v in model_summary[col].tolist():
            if pd.isna(v):
                out.append(0.0)
            else:
                out.append(float(v) * scale)
        return out

    charts["avg_ttft"] = _bar_chart(
        labels,
        col_vals("Avg TTFT (ms)"),
        "Average TTFT by Model",
        "ms",
        CHART_AVG_TTFT,
    )
    charts["p90_ttft"] = _bar_chart(
        labels,
        col_vals("P90 TTFT (ms)"),
        "P90 TTFT by Model",
        "ms",
        CHART_P90_TTFT,
    )
    charts["avg_latency"] = _bar_chart(
        labels,
        col_vals("Avg Total Time (ms)"),
        "Average Total Response Time by Model",
        "ms",
        CHART_AVG_LATENCY,
    )
    charts["tps"] = _bar_chart(
        labels,
        col_vals("Avg Tokens/sec"),
        "Tokens/sec by Model",
        "tokens/sec",
        CHART_TPS,
    )
    charts["success"] = _bar_chart(
        labels,
        col_vals("Success Rate", 100.0),
        "Success Rate by Model",
        "%",
        CHART_SUCCESS,
    )
    charts["accuracy"] = _bar_chart(
        labels,
        col_vals("Accuracy", 100.0),
        "Accuracy by Model",
        "%",
        CHART_ACCURACY,
    )

    # CueAI realtime subset charts
    if not raw.empty and "category" in raw.columns:
        cue = raw[raw["category"] == "REALTIME_INTERVIEW"]
        if not cue.empty:
            rows = []
            for keys, g in cue.groupby(["provider", "model_name"], sort=False):
                ok = g[g["status"] == "SUCCESS"]
                tt = [float(x) for x in ok["ttft_ms"].dropna().tolist()]
                tot = [float(x) for x in ok["total_latency_ms"].dropna().tolist()]
                rows.append(
                    {
                        "provider": keys[0],
                        "model": keys[1],
                        "avg_ttft": (sum(tt) / len(tt)) if tt else None,
                        "avg_total": (sum(tot) / len(tot)) if tot else None,
                    }
                )
            cdf = pd.DataFrame(rows).dropna(subset=["avg_ttft"], how="all")
            if not cdf.empty:
                clabels = [
                    f"{r.provider}\n{_safe_label(str(r.model))}"
                    for r in cdf.itertuples()
                ]
                charts["cueai_ttft"] = _bar_chart(
                    clabels,
                    [float(x) if pd.notna(x) else 0.0 for x in cdf["avg_ttft"]],
                    "CueAI Real-Time Average TTFT",
                    "ms",
                    CHART_CUEAI_TTFT,
                )
                charts["cueai_latency"] = _bar_chart(
                    clabels,
                    [float(x) if pd.notna(x) else 0.0 for x in cdf["avg_total"]],
                    "CueAI Real-Time Average Total Latency",
                    "ms",
                    CHART_CUEAI_LATENCY,
                )

        cat_rows = []
        for cat, g in raw.groupby("category", sort=False):
            ok = g[g["status"] == "SUCCESS"]
            tot = [float(x) for x in ok["total_latency_ms"].dropna().tolist()]
            accs = [float(x) for x in ok["accuracy_score"].dropna().tolist()] if "accuracy_score" in ok else []
            cat_rows.append(
                {
                    "category": cat,
                    "latency": (sum(tot) / len(tot)) if tot else 0.0,
                    "accuracy": (sum(accs) / len(accs) * 100.0) if accs else 0.0,
                }
            )
        if cat_rows:
            cdf = pd.DataFrame(cat_rows)
            clabels = [_safe_label(str(c), 18) for c in cdf["category"].tolist()]
            charts["category_latency"] = _bar_chart(
                clabels,
                [float(x) for x in cdf["latency"]],
                "Average Total Latency by Category",
                "ms",
                CHART_CATEGORY_LATENCY,
            )
            charts["category_accuracy"] = _bar_chart(
                clabels,
                [float(x) for x in cdf["accuracy"]],
                "Accuracy by Category",
                "%",
                CHART_CATEGORY_ACCURACY,
            )

        q_rows = []
        if "question_id" in raw.columns:
            for qid, g in raw.groupby("question_id", sort=False):
                ok = g[g["status"] == "SUCCESS"]
                tot = [float(x) for x in ok["total_latency_ms"].dropna().tolist()]
                if tot:
                    q_rows.append((str(qid), sum(tot) / len(tot)))
            q_rows = q_rows[:24]
            if q_rows:
                charts["question_latency"] = _bar_chart(
                    [q for q, _ in q_rows],
                    [v for _, v in q_rows],
                    "Question-Level Average Total Latency (first 24)",
                    "ms",
                    CHART_QUESTION_LATENCY,
                )
    return charts


def _write_df_sheet(ws, df: pd.DataFrame, title: str) -> None:
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    start = 3
    if df is None or df.empty:
        ws.cell(row=start, column=1, value="(no data)")
        return
    for r_idx, row in enumerate(
        dataframe_to_rows(df, index=False, header=True), start=start
    ):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == start:
                cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=False, vertical="center")
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width


def _attach_charts(ws, chart_paths: dict[str, Optional[Path]], start_row: int) -> None:
    row = start_row
    ws.cell(row=row, column=1, value="Charts").font = Font(bold=True, size=12)
    row += 2
    order = (
        "avg_ttft",
        "p90_ttft",
        "avg_latency",
        "tps",
        "success",
        "accuracy",
        "cueai_ttft",
        "cueai_latency",
        "category_latency",
        "category_accuracy",
        "question_latency",
    )
    for key in order:
        p = chart_paths.get(key)
        if p and Path(p).exists():
            img = XLImage(str(p))
            img.width = 720
            img.height = 320
            ws.add_image(img, f"A{row}")
            row += 18


def save_excel_report(
    *,
    raw: pd.DataFrame,
    model_summary: pd.DataFrame,
    question_summary: pd.DataFrame,
    chart_paths: dict[str, Optional[Path]],
    path: Path = EXCEL_PATH,
    category_summary: Optional[pd.DataFrame] = None,
) -> Path:
    ensure_results_dir()
    wb = Workbook()

    # 1. Executive Summary
    ws_exec = wb.active
    exec_cols = [
        "Provider",
        "Model",
        "Requests",
        "Success Rate",
        "Avg TTFT (ms)",
        "P90 TTFT (ms)",
        "Avg Total Time (ms)",
        "Avg Tokens/sec",
        "Accuracy",
        "Errors",
        "Timeouts",
    ]
    exec_df = (
        model_summary[[c for c in exec_cols if c in model_summary.columns]].copy()
        if not model_summary.empty
        else pd.DataFrame(columns=exec_cols)
    )
    if not exec_df.empty:
        exec_df = exec_df.sort_values(
            by=["Avg TTFT (ms)", "Avg Total Time (ms)"],
            ascending=[True, True],
            na_position="last",
        )
    _write_df_sheet(ws_exec, exec_df, "Executive Summary")
    _attach_charts(ws_exec, chart_paths, start_row=(len(exec_df) + 6))

    # 2. Model Summary
    ws_ms = wb.create_sheet()
    _write_df_sheet(ws_ms, model_summary, "Model Summary")

    ws_cat = wb.create_sheet()
    _write_df_sheet(
        ws_cat,
        category_summary if category_summary is not None else pd.DataFrame(),
        "Category Comparison",
    )

    ws_qc = wb.create_sheet()
    _write_df_sheet(ws_qc, question_summary, "Question Comparison")

    category_sheets = (
        ("Technical", "TECHNICAL"),
        ("Aptitude", "APTITUDE"),
        ("Logical Reasoning", "LOGICAL_REASONING"),
        ("Coding", "CODING"),
        ("Debugging", "DEBUGGING"),
        ("SQL", "SQL_DATABASE"),
        ("Computer Science", "COMPUTER_SCIENCE"),
        ("Scenarios", "SCENARIO"),
        ("Real-Time", "REALTIME_INTERVIEW"),
        ("HR", "BEHAVIORAL_HR"),
        ("Short Answers", "SHORT_ANSWER"),
        ("Follow-Up", "FOLLOW_UP"),
    )
    for title, cat in category_sheets:
        subset = (
            question_summary[question_summary["category"] == cat].copy()
            if not question_summary.empty and "category" in question_summary.columns
            else pd.DataFrame()
        )
        if subset.empty and not raw.empty and "category" in raw.columns:
            subset = raw[raw["category"] == cat].copy()
        ws = wb.create_sheet()
        _write_df_sheet(ws, subset, title)

    # 4. Speed Results
    speed_cols = [
        c
        for c in [
            "Provider",
            "Model",
            "Avg TTFT (ms)",
            "Median TTFT (ms)",
            "P90 TTFT (ms)",
            "P95 TTFT (ms)",
            "Min TTFT (ms)",
            "Max TTFT (ms)",
            "Avg Total Time (ms)",
            "Median Total Time (ms)",
            "P90 Total Time (ms)",
            "Avg Tokens/sec",
            "Median Tokens/sec",
        ]
        if c in model_summary.columns
    ]
    ws_speed = wb.create_sheet()
    _write_df_sheet(
        ws_speed,
        model_summary[speed_cols] if speed_cols else model_summary,
        "Speed Results",
    )

    # 5. Accuracy Results
    acc_cols = [
        c
        for c in [
            "Provider",
            "Model",
            "Accuracy",
            "Success Rate",
            "Requests",
            "Successful",
            "Failed",
        ]
        if c in model_summary.columns
    ]
    ws_acc = wb.create_sheet()
    _write_df_sheet(
        ws_acc,
        model_summary[acc_cols] if acc_cols else model_summary,
        "Accuracy Results",
    )

    # 6. CueAI Real-Time Results
    cue_raw = (
        raw[raw["category"] == "REALTIME_INTERVIEW"].copy()
        if not raw.empty and "category" in raw.columns
        else pd.DataFrame()
    )
    cue_q = (
        question_summary[question_summary["category"] == "REALTIME_INTERVIEW"].copy()
        if not question_summary.empty and "category" in question_summary.columns
        else pd.DataFrame()
    )
    ws_cue = wb.create_sheet()
    _write_df_sheet(
        ws_cue,
        cue_q if not cue_q.empty else cue_raw,
        "CueAI Real-Time Results",
    )

    # 7. Raw Results
    ws_raw = wb.create_sheet()
    _write_df_sheet(ws_raw, raw, "Raw Results")

    # 8. Errors
    if not raw.empty and "status" in raw.columns:
        errors = raw[raw["status"] != "SUCCESS"].copy()
    else:
        errors = pd.DataFrame()
    ws_err = wb.create_sheet()
    _write_df_sheet(ws_err, errors, "Errors")

    wb.save(path)
    return path


# Back-compat wrapper
def save_excel(
    raw: pd.DataFrame,
    summary: pd.DataFrame,
    chart_paths: dict[str, Optional[Path]],
    path: Path = EXCEL_PATH,
) -> Path:
    qsum = pd.DataFrame()
    return save_excel_report(
        raw=raw,
        model_summary=summary,
        question_summary=qsum,
        chart_paths=chart_paths,
        path=path,
    )


def print_final_comparison(model_summary: pd.DataFrame) -> None:
    print("\n" + "=" * 88)
    print("FINAL MODEL COMPARISON")
    print("=" * 88)
    if model_summary.empty:
        print("No summary rows.")
        print("=" * 88)
        return

    header = (
        f"{'Model':<28} {'Avg TTFT':>10} {'P90 TTFT':>10} "
        f"{'Avg Total':>10} {'Tok/s':>9} {'Acc':>7} {'OK%':>7}"
    )
    print(header)
    print("-" * 88)
    view = model_summary.copy()
    if "Avg TTFT (ms)" in view.columns:
        view = view.sort_values(
            by=["Avg TTFT (ms)"], ascending=True, na_position="last"
        )
    for _, row in view.iterrows():
        name = f"{row['Provider']}/{row['Model']}"
        if len(name) > 28:
            name = name[:27] + "."
        avg_ttft_v = row.get("Avg TTFT (ms)")
        p90_v = row.get("P90 TTFT (ms)")
        tot_v = row.get("Avg Total Time (ms)")
        tps_v = row.get("Avg Tokens/sec")
        acc_v = row.get("Accuracy")
        ok_v = row.get("Success Rate")
        avg_ttft_s = f"{avg_ttft_v:.0f}ms" if pd.notna(avg_ttft_v) else "-"
        p90_s = f"{p90_v:.0f}ms" if pd.notna(p90_v) else "-"
        tot_s = f"{tot_v:.0f}ms" if pd.notna(tot_v) else "-"
        tps_s = f"{tps_v:.1f}" if pd.notna(tps_v) else "-"
        acc_s = f"{100 * float(acc_v):.0f}%" if pd.notna(acc_v) else "-"
        ok_s = f"{100 * float(ok_v):.0f}%" if pd.notna(ok_v) else "-"
        print(
            f"{name:<28} {avg_ttft_s:>10} {p90_s:>10} {tot_s:>10} "
            f"{tps_s:>9} {acc_s:>7} {ok_s:>7}"
        )
    print("=" * 88)


def print_terminal_summary(summary: pd.DataFrame) -> None:
    print_final_comparison(summary)
